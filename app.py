import streamlit as st
from gigachat import GigaChat
import os
import io
import re

from docx import Document
import PyPDF2
from pptx import Presentation
import pymupdf  # современная замена fitz
from PIL import Image
import pytesseract

st.set_page_config(page_title="ИИ-Анализатор договоров", page_icon="🤖", layout="wide")

st.title("🤖 ИИ-Анализатор договоров")
st.markdown("Автоматический анализ договоров с помощью искусственного интеллекта")
st.markdown("📄 **Загрузите договоры** → 🤖 **ИИ найдёт риски** → 📊 **Получите отчёты**")


# ============================================================
# 1. АВТОПОДГРУЗКА ТОКЕНА ИЗ SECRETS
# ============================================================
def get_credentials():
    try:
        for key in ("GIGACHAT_CREDENTIALS", "GIGACHAT_TOKEN", "GIGACHAT_ACCESS_TOKEN"):
            if key in st.secrets:
                return str(st.secrets[key]).strip()
    except Exception:
        pass
    for var in ("GIGACHAT_CREDENTIALS", "GIGACHAT_TOKEN", "GIGACHAT_ACCESS_TOKEN"):
        val = os.getenv(var)
        if val:
            return val.strip()
    return None


# ============================================================
# 2. САЙДБАР
# ============================================================
with st.sidebar:
    st.header("⚙️ Настройки")
    creds = get_credentials()
    if creds:
        st.success("✅ Токен GigaChat загружен автоматически")
    else:
        st.warning("⚠️ Токен не найден. Добавьте GIGACHAT_CREDENTIALS в Settings → Secrets")
    
    st.markdown("---")
    model_name = st.selectbox(
        "Модель нейросети (код сам подберёт рабочую)",
        ["GigaChat-2-Pro", "GigaChat-2-Max", "GigaChat-Max", "GigaChat-Pro", "GigaChat"],
        index=0,
    )
    st.session_state["giga_model"] = model_name


# ============================================================
# 3. ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ ФАЙЛОВ
# ============================================================
def extract_text_from_docx(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join([p.text for p in doc.paragraphs])


def extract_text_from_pdf(file_bytes):
    text = ""
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            text += page.get_text() + "\n"
        if len(text.strip()) > 20:
            return text
    except Exception:
        pass
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            text += page.extract_text() + "\n"
    except Exception:
        pass
    return text


def extract_text_from_pptx(file_bytes):
    prs = Presentation(io.BytesIO(file_bytes))
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                lines.append(shape.text)
    return "\n".join(lines)


def extract_text(file):
    file_bytes = file.read()
    filename = file.name.lower()
    if filename.endswith('.docx'):
        return extract_text_from_docx(file_bytes)
    elif filename.endswith('.pdf'):
        return extract_text_from_pdf(file_bytes)
    elif filename.endswith('.pptx'):
        return extract_text_from_pptx(file_bytes)
    elif filename.endswith(('.png', '.jpg', '.jpeg', '.tiff')):
        try:
            return pytesseract.image_to_string(Image.open(io.BytesIO(file_bytes)), lang='rus+eng')
        except Exception:
            return ""
    else:
        try:
            return file_bytes.decode('utf-8', errors='ignore')
        except Exception:
            return ""


# ============================================================
# 4. АНАЛИЗ ДОГОВОРА ЧЕРЕЗ GIGACHAT (С АВТОПОДБОРОМ МОДЕЛИ)
# ============================================================
def analyze_contract(text):
    credentials = get_credentials()
    if not credentials:
        raise ValueError("Не найден токен GigaChat. Добавьте GIGACHAT_CREDENTIALS в Settings → Secrets.")

    selected = st.session_state.get("giga_model", "GigaChat-2-Pro")
    
    # Все возможные актуальные модели Сбера (2025-2026)
    fallback_models = [
        selected,
        "GigaChat-2-Pro",
        "GigaChat-2-Max",
        "GigaChat-2",
        "GigaChat-Max",
        "GigaChat-Pro",
        "GigaChat",
        "GigaChat-Lite",
        "GigaChat-Ultra"
    ]
    # Убираем дубликаты, сохраняя порядок
    models_to_try = list(dict.fromkeys(fallback_models))
    
    prompt = f"""Проанализируй текст договора как профессиональный юрист и верни данные СТРОГО в следующем формате (каждый параметр с новой строки):
1. ТИП ДОГОВОРА: [тип договора]
2. СУБЪЕКТНЫЙ СОСТАВ: [все стороны договора с их ролями (Заказчик/Исполнитель/Поставщик и т.д.)]
3. ПРЕДМЕТ ДОГОВОРА: [краткое описание предмета]
4. СУММА ДОГОВОРА: [сумма и валюта]
5. СРОК ДЕЙСТВИЯ: [срок действия или дата окончания]
6. ПОРЯДОК ОПЛАТЫ: [условия и сроки оплаты]
7. ОТВЕТСТВЕННОСТЬ СТОРОН: [пени, штрафы, неустойки]
8. УСЛОВИЯ РАСТОРЖЕНИЯ: [порядок расторжения]
9. ИНН СТОРОН: [ИНН всех сторон или "отсутствует"]
10. ЮРИДИЧЕСКИЕ РИСКИ: [основные риски, каждый через точку с запятой]

ВАЖНО: отвечай только по фактам из текста. Если информации нет — пиши "не указано в тексте". В поле "Субъектный состав" пиши ТОЛЬКО названия компаний и их роли, НЕ пиши туда условия ответственности.

ТЕКСТ ДОГОВОРА:
{text[:15000]}"""

    last_error = None
    for model in models_to_try:
        try:
            giga = GigaChat(
                credentials=credentials,
                scope="GIGACHAT_API_PERS",
                model=model,
                verify_ssl_certs=False
            )
            response = giga.chat(prompt)
            st.session_state["working_model"] = model
            return response.choices[0].message.content
        except Exception as e:
            last_error = str(e)
            if "No such model" in last_error or "404" in last_error:
                continue
            else:
                raise
    
    raise ValueError(f"Ни одна модель не подошла. Ошибка: {last_error}")


# ============================================================
# 5. ПАРСИНГ ОТВЕТА ИИ (10 ПАРАМЕТРОВ)
# ============================================================
def empty_parsed():
    return {
        "Тип договора": "не определён",
        "Субъектный состав": "не определён",
        "Предмет договора": "не указан",
        "Сумма договора": "не указана",
        "Срок действия": "не указан",
        "Порядок оплаты": "не указан",
        "Ответственность сторон": "не указана",
        "Условия расторжения": "не указаны",
        "ИНН сторон": "отсутствует",
        "Юридические риски": "не найдены"
    }


def parse_response(ai_text):
    data = empty_parsed()
    if not ai_text:
        return data
    
    for line in ai_text.split("\n"):
        line = line.strip().lstrip("-•* ").replace("**", "")
        line = re.sub(r'^\d+[.)]\s*', '', line)
        if ":" not in line:
            continue
        key, val = line.split(":", 1)
        key = key.strip().upper()
        val = val.strip()
        if not val:
            continue
        
        if "ТИП" in key:
            data["Тип договора"] = val
        elif "СУБЪЕКТ" in key or ("СОСТАВ" in key and "ИНН" not in key):
            data["Субъектный состав"] = val
        elif "ПРЕДМЕТ" in key:
            data["Предмет договора"] = val
        elif "СУММА" in key or ("ЦЕНА" in key and "РАСТОРЖ" not in key):
            data["Сумма договора"] = val
        elif "СРОК" in key:
            data["Срок действия"] = val
        elif "ОПЛАТ" in key:
            data["Порядок оплаты"] = val
        elif "ОТВЕТСТВЕННОСТЬ" in key or "ПЕН" in key or "ШТРАФ" in key or "НЕУСТОЙК" in key:
            data["Ответственность сторон"] = val
        elif "РАСТОРЖ" in key:
            data["Условия расторжения"] = val
        elif "ИНН" in key:
            data["ИНН сторон"] = val
        elif "РИСК" in key:
            data["Юридические риски"] = val
    return data


# ============================================================
# 6. ГЕНЕРАЦИЯ WORD-ОТЧЁТА
# ============================================================
def generate_docx_report(fname, parsed, ai_response):
    doc = Document()
    doc.add_heading(f'Отчёт по анализу: {fname}', 0)
    doc.add_paragraph('Подготовлено: ИИ-Анализатор договоров')
    
    doc.add_heading('📋 Параметры анализа (10 пунктов)', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Параметр'
    hdr_cells[1].text = 'Значение'
    
    for key, val in parsed.items():
        if key != "Юридические риски":
            row_cells = table.add_row().cells
            row_cells[0].text = key
            row_cells[1].text = str(val)
    
    doc.add_heading('⚠️ Юридические риски', level=1)
    doc.add_paragraph(parsed.get("Юридические риски", "не найдены"))
    
    doc.add_heading('🤖 Полный ИИ-анализ', level=1)
    doc.add_paragraph(ai_response if ai_response else "Анализ не выполнен")
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# 7. ГЕНЕРАЦИЯ ПРЕЗЕНТАЦИИ PPTX
# ============================================================
def generate_pptx_report(fname, parsed, ai_response):
    prs = Presentation()
    
    # Слайд 1: Титульный
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"Анализ договора\n{fname}"
    slide.placeholders[1].text = "Подготовлено ИИ-Анализатором (10 параметров)"
    
    # Слайд 2: Основные параметры
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "📋 Основные параметры"
    tf = slide.placeholders[1].text_frame
    for i, key in enumerate(["Тип договора", "Субъектный состав", "Предмет договора", "Сумма договора", "Срок действия"]):
        if i == 0:
            tf.text = f"{key}: {parsed.get(key, 'не указано')}"
        else:
            p = tf.add_paragraph()
            p.text = f"{key}: {parsed.get(key, 'не указано')}"
    
    # Слайд 3: Условия
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "📄 Условия договора"
    tf = slide.placeholders[1].text_frame
    for i, key in enumerate(["Порядок оплаты", "Ответственность сторон", "Условия расторжения", "ИНН сторон"]):
        if i == 0:
            tf.text = f"{key}: {parsed.get(key, 'не указано')}"
        else:
            p = tf.add_paragraph()
            p.text = f"{key}: {parsed.get(key, 'не указано')}"
    
    # Слайд 4: Риски
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "⚠️ Юридические риски"
    slide.placeholders[1].text_frame.text = parsed.get("Юридические риски", "не найдены")
    
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# 8. ОСНОВНОЙ ИНТЕРФЕЙС
# ============================================================
st.markdown("---")
st.subheader("📤 Загрузка договоров")
uploaded_files = st.file_uploader(
    "Перетащите файлы",
    type=['docx', 'pdf', 'pptx', 'txt', 'png', 'jpg', 'jpeg'],
    accept_multiple_files=True
)

if uploaded_files:
    st.success(f"✅ Загружено файлов: {len(uploaded_files)}")
    results = []
    processed = 0

    for file in uploaded_files:
        st.markdown(f"📄 **{file.name}**")
        text = extract_text(file)
        if not text.strip():
            st.warning("⚠️ Не удалось извлечь текст.")
            continue

        model_hint = st.session_state.get('working_model', 'подбираю модель...')
        with st.spinner(f"🤖 ({model_hint}) анализирует документ по 10 параметрам..."):
            ai_response = None
            ai_error = None
            try:
                ai_response = analyze_contract(text)
            except Exception as e:
                ai_error = str(e)

        parsed = parse_response(ai_response)
        results.append((file.name, parsed, ai_response, ai_error))
        processed += 1

    if processed:
        st.balloons()
        st.success(f"🎉 Готово! Проанализировано: {processed} договоров")
        st.markdown("---")
        st.subheader("📊 Результаты анализа")

        for fname, parsed, ai_response, ai_error in results:
            st.markdown(f"#### 📄 {fname}")
            
            if ai_error:
                st.error(f"❌ Ошибка: {ai_error}")
                st.markdown("---")
                continue
            
            # Таблица с 9 параметрами
            param_df = [[k, v] for k, v in parsed.items() if k != "Юридические риски"]
            st.table(param_df)
            
            with st.expander("⚠️ Юридические риски", expanded=True):
                st.warning(parsed.get("Юридические риски", "не найдены"))
            with st.expander("🤖 Полный ответ ИИ"):
                st.info(ai_response)
            
            st.markdown("##### 📥 Скачать результаты:")
            col1, col2, col3 = st.columns(3)
            
            with col1:
                docx_buffer = generate_docx_report(fname, parsed, ai_response)
                st.download_button(
                    label="📄 Отчёт (Word)",
                    data=docx_buffer,
                    file_name=f"Report_{fname}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )
            with col2:
                pptx_buffer = generate_pptx_report(fname, parsed, ai_response)
                st.download_button(
                    label="📊 Презентация (PPTX)",
                    data=pptx_buffer,
                    file_name=f"Presentation_{fname}.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True
                )
            with col3:
                st.download_button(
                    label="📝 Текст (TXT)",
                    data=f"Отчёт по {fname}\n\n{ai_response or ''}".encode('utf-8'),
                    file_name=f"Report_{fname}.txt",
                    mime="text/plain",
                    use_container_width=True
                )
            st.markdown("---")
